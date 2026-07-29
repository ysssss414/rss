from __future__ import annotations

import json
import unittest
from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from three_board_rsi_entry.config import StrategyConfig
from three_board_rsi_entry.input_excel import InputData, create_input_template
from three_board_rsi_entry.outputs import write_outputs
from three_board_rsi_entry.pipeline import AnalysisRun

from .helpers import TEST_TEMP_ROOT, replay_one, trading_days


class OutputTests(unittest.TestCase):
    def test_template_contains_required_sheets_and_fields(self):
        path = create_input_template(Path(TEST_TEMP_ROOT) / "template_case.xlsx")
        workbook = load_workbook(path, read_only=True)
        self.assertEqual(
            workbook.sheetnames, ["交易日确认", "三连板股票"]
        )
        confirmation_headers = [
            cell.value for cell in next(workbook["交易日确认"].iter_rows())
        ]
        board_headers = [
            cell.value for cell in next(workbook["三连板股票"].iter_rows())
        ]
        self.assertEqual(
            confirmation_headers, ["trade_date", "input_complete", "note"]
        )
        self.assertIn("board_count", board_headers)

    def test_all_required_output_files_and_excel_sheets_are_written(self):
        days = trading_days(4)
        result = replay_one(days=days, as_of_index=2)
        input_data = InputData(
            confirmations=pd.DataFrame(),
            boards=pd.DataFrame(),
            missing_confirmation_dates=[],
            warnings=[],
            data_complete=True,
        )
        root = Path(TEST_TEMP_ROOT)
        input_path = root / "output_input_case.xlsx"
        input_path.write_bytes(b"test")
        run = AnalysisRun(
            result=result,
            input_data=input_data,
            checks=pd.DataFrame(
                [{"check_name": "test", "status": "PASS", "details": "ok"}]
            ),
            warnings=[],
            config=StrategyConfig(),
            start_date=days[0],
            as_of_date=days[2],
            input_path=input_path,
            input_hash="abc",
            market_source="TestProvider",
        )
        paths = write_outputs(run, root / "outputs_case")
        self.assertTrue(all(path.exists() for path in paths.values()))
        workbook = load_workbook(paths["candidate_status"], read_only=True)
        self.assertEqual(
            workbook.sheetnames,
            ["当前候选", "历史信号", "未准入连板", "运行检查", "参数"],
        )
        summary = json.loads(paths["run_summary"].read_text(encoding="utf-8"))
        self.assertEqual(summary["candidate_cycle_count"], 1)
        self.assertEqual(summary["run_as_of"], days[2].isoformat())

    def test_incomplete_input_warning_is_written_to_summary(self):
        days = trading_days(3)
        result = replay_one(days=days, as_of_index=1)
        warning = f"Missing completed trading-day confirmations: {days[1]}"
        input_data = InputData(
            confirmations=pd.DataFrame(),
            boards=pd.DataFrame(),
            missing_confirmation_dates=[days[1]],
            warnings=[warning],
            data_complete=False,
        )
        root = Path(TEST_TEMP_ROOT)
        input_path = root / "incomplete_input_case.xlsx"
        input_path.write_bytes(b"test")
        run = AnalysisRun(
            result=result,
            input_data=input_data,
            checks=pd.DataFrame(
                [
                    {
                        "check_name": "trading_day_confirmations",
                        "status": "WARN",
                        "details": warning,
                    }
                ]
            ),
            warnings=[warning],
            config=StrategyConfig(),
            start_date=days[0],
            as_of_date=days[1],
            input_path=input_path,
            input_hash="def",
            market_source="TestProvider",
        )
        paths = write_outputs(run, root / "incomplete_outputs_case")
        summary = json.loads(paths["run_summary"].read_text(encoding="utf-8"))
        self.assertEqual(summary["missing_confirmation_dates"], [str(days[1])])
        self.assertEqual(summary["warnings"], [warning])


if __name__ == "__main__":
    unittest.main()
