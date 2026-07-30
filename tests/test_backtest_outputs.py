from __future__ import annotations

import json
import unittest
from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from three_board_rsi_entry.backtest_outputs import write_backtest_outputs
from three_board_rsi_entry.config import StrategyConfig
from three_board_rsi_entry.input_excel import InputData
from three_board_rsi_entry.pipeline import AnalysisRun

from .backtest_helpers import run_backtest
from .helpers import TEST_TEMP_ROOT, replay_one, trading_days


class BacktestOutputTests(unittest.TestCase):
    def _analysis(self, days: list[date]) -> AnalysisRun:
        input_path = Path(TEST_TEMP_ROOT) / "backtest_input.xlsx"
        input_path.write_bytes(b"test")
        return AnalysisRun(
            result=replay_one(days=days, as_of_index=len(days) - 1),
            input_data=InputData(
                confirmations=pd.DataFrame(),
                boards=pd.DataFrame(),
                missing_confirmation_dates=[],
                warnings=[],
                data_complete=True,
            ),
            checks=pd.DataFrame(),
            warnings=[],
            config=StrategyConfig(),
            start_date=days[0],
            as_of_date=days[-1],
            input_path=input_path,
            input_hash="abc",
            market_source="TestProvider",
        )

    def test_all_backtest_outputs_and_workbook_sheets_are_written(self):
        days = trading_days(8)
        backtest = run_backtest(days=days)
        root = Path(TEST_TEMP_ROOT) / "backtest_outputs"
        paths = write_backtest_outputs(backtest, self._analysis(days), root)

        self.assertEqual(
            set(paths),
            {
                "backtest_events",
                "first_signal_events",
                "backtest_summary",
                "backtest_run_summary",
            },
        )
        self.assertTrue(all(path.exists() for path in paths.values()))
        workbook = load_workbook(paths["backtest_summary"], read_only=True)
        self.assertEqual(
            workbook.sheetnames,
            ["事件明细", "首信号明细", "信号汇总", "分层统计", "未完成样本", "参数"],
        )
        events = pd.read_csv(paths["backtest_events"])
        self.assertIn("entry_status", events)
        self.assertIn("holding_status_10d", events)

    def test_run_summary_counts_and_parameters_are_auditable(self):
        days = trading_days(8)
        backtest = run_backtest(days=days)
        root = Path(TEST_TEMP_ROOT) / "backtest_summary_case"
        paths = write_backtest_outputs(backtest, self._analysis(days), root)

        summary = json.loads(
            paths["backtest_run_summary"].read_text(encoding="utf-8")
        )
        self.assertEqual(summary["holding_periods"], [1, 3, 5, 10])
        self.assertEqual(summary["entry_price_rule"], "NEXT_TRADABLE_OPEN")
        self.assertEqual(summary["raw_signal_count"], 1)
        self.assertEqual(summary["completed_5d_count"], 1)
        self.assertEqual(summary["completed_10d_count"], 0)
        parameters = pd.read_excel(paths["backtest_summary"], sheet_name="参数")
        self.assertEqual(
            parameters.loc[
                parameters["parameter"] == "entry_price_rule", "value"
            ].iloc[0],
            "NEXT_TRADABLE_OPEN",
        )

    def test_repeated_output_is_byte_deterministic(self):
        days = trading_days(8)
        backtest = run_backtest(days=days)
        analysis = self._analysis(days)
        first = write_backtest_outputs(
            backtest, analysis, Path(TEST_TEMP_ROOT) / "backtest_deterministic_a"
        )
        second = write_backtest_outputs(
            backtest, analysis, Path(TEST_TEMP_ROOT) / "backtest_deterministic_b"
        )

        for key in first:
            self.assertEqual(first[key].read_bytes(), second[key].read_bytes())


if __name__ == "__main__":
    unittest.main()
