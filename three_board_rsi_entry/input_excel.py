from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Iterable

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


CONFIRMATION_SHEET = "交易日确认"
BOARD_SHEET = "三连板股票"
CONFIRMATION_COLUMNS = ["trade_date", "input_complete", "note"]
BOARD_COLUMNS = [
    "trade_date",
    "ts_code",
    "stock_name",
    "board_count",
    "limit_up_type",
    "note",
]


@dataclass(frozen=True)
class InputData:
    confirmations: pd.DataFrame
    boards: pd.DataFrame
    missing_confirmation_dates: list[date]
    warnings: list[str]
    data_complete: bool


def _format_template_sheet(worksheet, widths: list[int]) -> None:
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in worksheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = fill
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[chr(64 + index)].width = width


def create_input_template(output: Path | str) -> Path:
    output_path = Path(output)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    confirmations = workbook.active
    confirmations.title = CONFIRMATION_SHEET
    confirmations.append(CONFIRMATION_COLUMNS)
    _format_template_sheet(confirmations, [16, 18, 36])

    boards = workbook.create_sheet(BOARD_SHEET)
    boards.append(BOARD_COLUMNS)
    _format_template_sheet(boards, [16, 18, 18, 14, 16, 36])

    workbook.save(output_path)
    return output_path


def normalize_ts_code(value: object) -> str:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        raise ValueError("ts_code cannot be empty")
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip().upper()
    if text.endswith(".0") and text[:-2].isdigit():
        text = text[:-2]
    if text.isdigit() and len(text) <= 6:
        text = text.zfill(6)
    if not text:
        raise ValueError("ts_code cannot be empty")
    return text


def _normalize_date(value: object, field: str) -> date:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        raise ValueError(f"{field} cannot be empty")
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return pd.Timestamp(value).date()
    except (TypeError, ValueError) as exc:
        raise ValueError(f"Invalid {field}: {value!r}; expected YYYY-MM-DD") from exc


def _require_columns(frame: pd.DataFrame, required: Iterable[str], sheet: str) -> None:
    missing = sorted(set(required) - set(frame.columns))
    if missing:
        raise ValueError(f"Sheet '{sheet}' missing fields: {', '.join(missing)}")


def _next_day_map(trading_days: list[date]) -> dict[date, date]:
    return dict(zip(trading_days, trading_days[1:]))


def validate_board_sequences(boards: pd.DataFrame, trading_days: list[date]) -> None:
    next_day = _next_day_map(trading_days)
    for ts_code, rows in boards.groupby("ts_code", sort=False):
        ordered = rows.sort_values("trade_date", kind="stable")
        previous = None
        for row in ordered.itertuples(index=False):
            if previous is not None and next_day.get(previous.trade_date) == row.trade_date:
                expected = previous.board_count + 1
                if row.board_count != expected:
                    raise ValueError(
                        "Adjacent trading-day board_count must increase by 1: "
                        f"{ts_code} {row.trade_date} expected {expected}, "
                        f"got {row.board_count} (previous {previous.board_count})"
                    )
            previous = row


def load_input_workbook(
    path: Path | str,
    *,
    start_date: date,
    as_of_date: date,
    trading_days: Iterable[date],
    allow_incomplete: bool = False,
) -> InputData:
    input_path = Path(path)
    if start_date > as_of_date:
        raise ValueError("start_date cannot be later than as_of_date")
    try:
        sheets = pd.read_excel(
            input_path,
            sheet_name=[CONFIRMATION_SHEET, BOARD_SHEET],
            engine="openpyxl",
            dtype=object,
        )
    except FileNotFoundError as exc:
        raise ValueError(f"Input workbook not found: {input_path}") from exc
    except ValueError as exc:
        raise ValueError(
            f"Input workbook must contain sheets '{CONFIRMATION_SHEET}' and "
            f"'{BOARD_SHEET}': {exc}"
        ) from exc

    confirmations = sheets[CONFIRMATION_SHEET].dropna(how="all").copy()
    boards = sheets[BOARD_SHEET].dropna(how="all").copy()
    _require_columns(confirmations, CONFIRMATION_COLUMNS[:2], CONFIRMATION_SHEET)
    _require_columns(boards, ["trade_date", "ts_code", "board_count"], BOARD_SHEET)

    confirmations["trade_date"] = confirmations["trade_date"].map(
        lambda value: _normalize_date(value, "trade_date")
    )
    boards["trade_date"] = boards["trade_date"].map(
        lambda value: _normalize_date(value, "trade_date")
    )

    # Future rows are intentionally discarded before validation. Historical
    # replay must behave exactly as it would have on that as-of date.
    confirmations = confirmations[
        (confirmations["trade_date"] >= start_date)
        & (confirmations["trade_date"] <= as_of_date)
    ].copy()
    boards = boards[
        (boards["trade_date"] >= start_date) & (boards["trade_date"] <= as_of_date)
    ].copy()

    if confirmations["trade_date"].duplicated().any():
        duplicates = sorted(
            confirmations.loc[
                confirmations["trade_date"].duplicated(keep=False), "trade_date"
            ].unique()
        )
        raise ValueError(
            "Duplicate confirmation dates: "
            + ", ".join(value.isoformat() for value in duplicates)
        )

    complete = pd.to_numeric(confirmations["input_complete"], errors="coerce")
    invalid_complete = confirmations.loc[complete != 1, "trade_date"].tolist()
    if invalid_complete:
        raise ValueError(
            "input_complete must be 1 for every recorded confirmation date: "
            + ", ".join(value.isoformat() for value in invalid_complete)
        )
    confirmations["input_complete"] = complete.astype(int)

    boards["ts_code"] = boards["ts_code"].map(normalize_ts_code)
    board_numeric = pd.to_numeric(boards["board_count"], errors="coerce")
    integral = board_numeric.notna() & (board_numeric % 1 == 0)
    if not integral.all():
        bad_rows = (boards.index[~integral] + 2).tolist()
        raise ValueError(f"board_count must be an integer; worksheet rows: {bad_rows}")
    boards["board_count"] = board_numeric.astype(int)
    below_three = boards[boards["board_count"] < 3]
    if not below_three.empty:
        details = ", ".join(
            f"{row.ts_code} {row.trade_date}={row.board_count}"
            for row in below_three.itertuples(index=False)
        )
        raise ValueError(f"board_count must be at least 3: {details}")

    duplicated = boards.duplicated(["trade_date", "ts_code"], keep=False)
    if duplicated.any():
        details = ", ".join(
            f"{row.ts_code} {row.trade_date}"
            for row in boards.loc[duplicated].itertuples(index=False)
        )
        raise ValueError(f"Duplicate trade_date/ts_code rows: {details}")

    calendar = sorted(
        {
            value
            for value in trading_days
            if start_date <= value <= as_of_date
        }
    )
    calendar_set = set(calendar)
    non_trading = sorted(set(boards["trade_date"]) - calendar_set)
    if non_trading:
        raise ValueError(
            "Board records contain non-trading dates: "
            + ", ".join(value.isoformat() for value in non_trading)
        )

    confirmation_dates = set(confirmations["trade_date"])
    missing_dates = [value for value in calendar if value not in confirmation_dates]
    warnings: list[str] = []
    if missing_dates:
        message = (
            "Missing completed trading-day confirmations: "
            + ", ".join(value.isoformat() for value in missing_dates)
        )
        if not allow_incomplete:
            raise ValueError(message)
        warnings.append(message)

    validate_board_sequences(boards, calendar)
    for optional in ["stock_name", "limit_up_type", "note"]:
        if optional not in boards:
            boards[optional] = ""
        boards[optional] = boards[optional].fillna("").astype(str).str.strip()
    if "note" not in confirmations:
        confirmations["note"] = ""
    confirmations["note"] = confirmations["note"].fillna("").astype(str).str.strip()

    confirmations = confirmations.sort_values("trade_date", kind="stable").reset_index(
        drop=True
    )
    boards = boards.sort_values(["trade_date", "ts_code"], kind="stable").reset_index(
        drop=True
    )
    return InputData(
        confirmations=confirmations,
        boards=boards,
        missing_confirmation_dates=missing_dates,
        warnings=warnings,
        data_complete=not missing_dates,
    )
