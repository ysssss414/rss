from __future__ import annotations

import json
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

from .pipeline import AnalysisRun


DATE_COLUMNS = {
    "signal_date",
    "sequence_start_date",
    "qualified_date",
    "last_limit_up_date",
    "latest_trade_date",
    "first_ma5_touch_date",
    "md1_date",
    "below70_start_date",
    "md2_date",
    "cycle_expiry_date",
    "superseded_date",
    "run_as_of",
}


def _serializable_frame(frame: pd.DataFrame) -> pd.DataFrame:
    result = frame.copy()
    for column in DATE_COLUMNS & set(result.columns):
        result[column] = result[column].map(
            lambda value: value.isoformat()
            if isinstance(value, (date, datetime))
            else ("" if pd.isna(value) else value)
        )
    return result


def _format_workbook(path: Path, as_of_date: date) -> None:
    workbook = load_workbook(path)
    fill = PatternFill("solid", fgColor="1F4E78")
    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        if worksheet.max_row >= 1 and worksheet.max_column >= 1:
            worksheet.auto_filter.ref = worksheet.dimensions
        for cell in worksheet[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = fill
        for column_cells in worksheet.columns:
            values = [str(cell.value or "") for cell in column_cells[:200]]
            width = min(max(max(map(len, values), default=8) + 2, 10), 42)
            worksheet.column_dimensions[column_cells[0].column_letter].width = width
    fixed_timestamp = datetime.combine(as_of_date, datetime.min.time())
    workbook.properties.created = fixed_timestamp
    workbook.properties.modified = fixed_timestamp
    workbook.save(path)


def _normalize_xlsx_archive(path: Path, as_of_date: date) -> None:
    """Make ZIP metadata deterministic after openpyxl writes the workbook."""

    temporary = path.with_name(f".{path.name}.deterministic.tmp")
    with ZipFile(path, "r") as source, ZipFile(
        temporary, "w", compression=ZIP_DEFLATED, compresslevel=9
    ) as destination:
        for original in sorted(source.infolist(), key=lambda item: item.filename):
            normalized = ZipInfo(original.filename, date_time=(1980, 1, 1, 0, 0, 0))
            normalized.compress_type = ZIP_DEFLATED
            normalized.external_attr = original.external_attr
            normalized.create_system = original.create_system
            normalized.flag_bits = original.flag_bits
            payload = source.read(original.filename)
            if original.filename == "docProps/core.xml":
                fixed = f"{as_of_date.isoformat()}T00:00:00Z".encode("ascii")
                payload = re.sub(
                    rb"(<dcterms:modified\b[^>]*>).*?(</dcterms:modified>)",
                    lambda match: match.group(1) + fixed + match.group(2),
                    payload,
                )
            destination.writestr(normalized, payload)
    temporary.replace(path)


def _summary(run: AnalysisRun) -> dict[str, Any]:
    candidates = run.result.candidate_cycles
    signals = run.result.signals
    qualified_count = (
        int(candidates["qualified_date"].notna().sum()) if not candidates.empty else 0
    )
    md1_count = (
        int((signals["signal_type"] == "MD_1").sum()) if not signals.empty else 0
    )
    md2_count = (
        int((signals["signal_type"] == "MD_2").sum()) if not signals.empty else 0
    )
    return {
        "run_as_of": run.as_of_date.isoformat(),
        "start_date": run.start_date.isoformat(),
        "price_adjustment": run.config.price_adjustment,
        "rsi_period": run.config.rsi_period,
        "input_file": str(run.input_path),
        "input_file_hash": run.input_hash,
        "candidate_cycle_count": int(len(candidates)),
        "qualified_cycle_count": qualified_count,
        "md1_signal_count": md1_count,
        "md2_signal_count": md2_count,
        "missing_confirmation_dates": [
            value.isoformat()
            for value in run.input_data.missing_confirmation_dates
        ],
        "warnings": run.warnings,
    }


def write_outputs(run: AnalysisRun, output_dir: Path | str) -> dict[str, Path]:
    destination = Path(output_dir)
    destination.mkdir(parents=True, exist_ok=True)
    candidate_path = destination / "candidate_cycles.csv"
    signals_path = destination / "signals.csv"
    excel_path = destination / "candidate_status.xlsx"
    summary_path = destination / "run_summary.json"

    candidates = _serializable_frame(run.result.candidate_cycles)
    signals = _serializable_frame(run.result.signals)
    candidates.to_csv(
        candidate_path, index=False, lineterminator="\n", float_format="%.10g"
    )
    signals.to_csv(
        signals_path, index=False, lineterminator="\n", float_format="%.10g"
    )

    qualified = candidates[candidates["qualified_date"].astype(str) != ""].copy()
    unqualified = candidates[candidates["qualified_date"].astype(str) == ""].copy()
    parameters = pd.DataFrame(
        [
            {"parameter": key, "value": value}
            for key, value in run.config.as_dict().items()
        ]
        + [
            {"parameter": "run_as_of", "value": run.as_of_date.isoformat()},
            {"parameter": "start_date", "value": run.start_date.isoformat()},
            {"parameter": "market_source", "value": run.market_source},
        ]
    )
    with pd.ExcelWriter(excel_path, engine="openpyxl") as writer:
        qualified.to_excel(writer, sheet_name="当前候选", index=False)
        signals.to_excel(writer, sheet_name="历史信号", index=False)
        unqualified.to_excel(writer, sheet_name="未准入连板", index=False)
        run.checks.to_excel(writer, sheet_name="运行检查", index=False)
        parameters.to_excel(writer, sheet_name="参数", index=False)
    _format_workbook(excel_path, run.as_of_date)
    _normalize_xlsx_archive(excel_path, run.as_of_date)

    summary_path.write_text(
        json.dumps(_summary(run), ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    return {
        "candidate_cycles": candidate_path,
        "signals": signals_path,
        "candidate_status": excel_path,
        "run_summary": summary_path,
    }
